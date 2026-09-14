"""이미지 검토 중 작성한 불량 표시와 메모의 저장·Excel 추출."""

from __future__ import annotations

import hashlib
import json
import os
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PyQt6.QtCore import QObject, pyqtSignal

from models import ExtractedImage, ROLE_ORDER


ROLE_DISPLAY_NAMES = {
    "ref": "Excel Ref",
    "a": "Excel 비교A",
    "b": "Excel 비교B",
    "c": "Excel 비교C",
}


@dataclass
class ReviewAnnotation:
    """한 Excel 원본 이미지에 연결된 사용자 검토 기록."""

    key: str
    role: str
    workbook_path: str
    workbook_name: str
    sheet_name: str
    sheet_index: int
    cell_address: str
    merged_range: str = ""
    anchor_occurrence: int = 1
    anchor_count: int = 1
    is_defect: bool = False
    note: str = ""
    updated_at: str = ""

    @property
    def should_keep(self) -> bool:
        return self.is_defect or bool(self.note.strip())

    @property
    def record_type(self) -> str:
        if self.is_defect and self.note.strip():
            return "불량 + 메모"
        if self.is_defect:
            return "불량"
        return "메모"


def default_review_store_path() -> Path:
    base = os.environ.get("LOCALAPPDATA")
    if not base:
        base = str(Path.home() / "AppData" / "Local")
    return Path(base) / "excel-image-inspector" / "review_annotations.json"


def workbook_session_key(workbook_paths: Dict[str, str], mode: str) -> str:
    """경로와 파일 상태가 같은 Excel 조합에 대해 안정적인 세션 키를 만든다."""
    sources = []
    for role in ROLE_ORDER:
        raw_path = workbook_paths.get(role, "")
        if not raw_path:
            continue
        path = os.path.normcase(os.path.abspath(raw_path))
        try:
            stat = os.stat(path)
            file_state = [stat.st_size, stat.st_mtime_ns]
        except OSError:
            file_state = [None, None]
        sources.append([role, path, *file_state])
    raw = json.dumps([mode, sources], ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def image_annotation_key(role: str, image: ExtractedImage) -> str:
    identity = image.source_image_id or "|".join(
        (
            str(image.sheet_index),
            image.source_sheet_id or image.sheet_name,
            str(image.anchor_row),
            str(image.anchor_col),
            str(image.anchor_occurrence),
        )
    )
    return f"{role}|{identity}"


class ReviewAnnotationStore(QObject):
    """현재 Excel 조합의 검토 기록을 즉시 JSON으로 자동 저장한다."""

    changed = pyqtSignal()

    def __init__(
        self,
        workbook_paths: Dict[str, str],
        mode: str,
        *,
        storage_path: Optional[Path | str] = None,
        parent=None,
    ) -> None:
        super().__init__(parent)
        self.workbook_paths = dict(workbook_paths)
        self.mode = mode
        self.storage_path = Path(storage_path or default_review_store_path())
        self.session_key = workbook_session_key(self.workbook_paths, mode)
        self._document = self._load_document()
        sessions = self._document.setdefault("sessions", {})
        session = sessions.setdefault(
            self.session_key,
            {
                "mode": mode,
                "workbook_paths": self.workbook_paths,
                "records": {},
            },
        )
        self._records: Dict[str, dict] = session.setdefault("records", {})

    def _load_document(self) -> dict:
        try:
            with self.storage_path.open("r", encoding="utf-8") as stream:
                payload = json.load(stream)
            if isinstance(payload, dict) and isinstance(payload.get("sessions"), dict):
                return payload
        except (OSError, ValueError, TypeError):
            pass
        return {"version": 1, "sessions": {}}

    def _save(self) -> None:
        self.storage_path.parent.mkdir(parents=True, exist_ok=True)
        temporary = self.storage_path.with_suffix(self.storage_path.suffix + ".tmp")
        with temporary.open("w", encoding="utf-8") as stream:
            json.dump(self._document, stream, ensure_ascii=False, indent=2)
        os.replace(temporary, self.storage_path)

    def annotation_for(
        self, role: str, image: Optional[ExtractedImage]
    ) -> Optional[ReviewAnnotation]:
        if image is None or image.is_null:
            return None
        key = image_annotation_key(role, image)
        stored = self._records.get(key)
        if stored is not None:
            return ReviewAnnotation(**stored)
        path = self.workbook_paths.get(role, "")
        return ReviewAnnotation(
            key=key,
            role=role,
            workbook_path=path,
            workbook_name=os.path.basename(path) if path else "-",
            sheet_name=image.sheet_name,
            sheet_index=image.sheet_index,
            cell_address=image.cell_address,
            merged_range=image.merged_range,
            anchor_occurrence=image.anchor_occurrence,
            anchor_count=image.anchor_count,
        )

    def update(
        self,
        role: str,
        image: Optional[ExtractedImage],
        *,
        is_defect: Optional[bool] = None,
        note: Optional[str] = None,
    ) -> Optional[ReviewAnnotation]:
        annotation = self.annotation_for(role, image)
        if annotation is None:
            return None
        if is_defect is not None:
            annotation.is_defect = bool(is_defect)
        if note is not None:
            annotation.note = str(note)
        annotation.updated_at = datetime.now().astimezone().isoformat(timespec="seconds")
        if annotation.should_keep:
            self._records[annotation.key] = asdict(annotation)
        else:
            self._records.pop(annotation.key, None)
        self._save()
        self.changed.emit()
        return annotation

    def records(self) -> list[ReviewAnnotation]:
        records = [ReviewAnnotation(**payload) for payload in self._records.values()]
        return sorted(
            (record for record in records if record.should_keep),
            key=lambda record: (
                record.sheet_index,
                record.sheet_name,
                record.cell_address,
                ROLE_ORDER.index(record.role) if record.role in ROLE_ORDER else 99,
                record.anchor_occurrence,
            ),
        )

    @property
    def defect_count(self) -> int:
        return sum(record.is_defect for record in self.records())

    @property
    def memo_only_count(self) -> int:
        return sum(
            bool(record.note.strip()) and not record.is_defect
            for record in self.records()
        )


def export_annotations_xlsx(
    output_path: Path | str,
    records: Iterable[ReviewAnnotation],
) -> int:
    """불량 또는 메모가 있는 기록만 사용자가 읽기 쉬운 Excel로 만든다."""
    kept = [record for record in records if record.should_keep]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "검토 기록"
    headers = [
        "번호",
        "기록 유형",
        "불량 여부",
        "대상",
        "파일명",
        "파일 경로",
        "시트",
        "셀 위치",
        "병합 영역",
        "중복 순번",
        "메모",
        "기록 시간",
    ]
    sheet.append(headers)
    for index, record in enumerate(kept, start=1):
        duplicate = (
            f"{record.anchor_occurrence}/{record.anchor_count}"
            if record.anchor_count > 1
            else ""
        )
        sheet.append(
            [
                index,
                record.record_type,
                "예" if record.is_defect else "아니오",
                ROLE_DISPLAY_NAMES.get(record.role, record.role),
                record.workbook_name,
                record.workbook_path,
                record.sheet_name,
                record.cell_address,
                record.merged_range,
                duplicate,
                record.note,
                record.updated_at,
            ]
        )
        sheet.cell(index + 1, 11).data_type = "s"

    header_fill = PatternFill("solid", fgColor="1E293B")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = [7, 14, 11, 14, 34, 52, 24, 14, 22, 12, 60, 25]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:L{max(1, len(kept) + 1)}"
    sheet.row_dimensions[1].height = 24
    for index in range(2, len(kept) + 2):
        sheet.row_dimensions[index].height = 42
    workbook.save(str(output_path))
    return len(kept)
