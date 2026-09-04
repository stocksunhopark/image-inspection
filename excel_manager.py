"""openpyxl 기반 Excel 이미지 위치 추출과 지연 미리보기 로딩."""

import os
import zipfile
from typing import Dict, List, Optional, Tuple
from xml.etree.ElementTree import ParseError, fromstring

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image

from models import ExtractedImage

Position = Tuple[int, int]


def _xml_local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _sheet_names_from_workbook_xml(path: str) -> Optional[List[str]]:
    try:
        with zipfile.ZipFile(path, "r") as archive:
            xml = archive.read("xl/workbook.xml")
    except (OSError, KeyError, zipfile.BadZipFile):
        return None
    try:
        root = fromstring(xml)
    except ParseError:
        return None
    sheets_element = next(
        (child for child in list(root) if _xml_local_name(child.tag) == "sheets"),
        None,
    )
    if sheets_element is None:
        return None
    names = [
        element.get("name")
        for element in list(sheets_element)
        if _xml_local_name(element.tag) == "sheet" and element.get("name")
    ]
    return names or None


def read_sheet_index_names(path: str) -> List[Tuple[int, str]]:
    """이미지나 스타일을 열지 않고 ``(1-based 인덱스, 시트명)``을 반환한다."""
    titles = _sheet_names_from_workbook_xml(path)
    if titles is None:
        workbook = load_workbook(path, read_only=True)
        try:
            titles = list(workbook.sheetnames)
        finally:
            workbook.close()
    return list(enumerate(titles, start=1))


def find_merged_range(worksheet, row_1based: int, col_1based: int) -> str:
    for merged in worksheet.merged_cells.ranges:
        if (
            merged.min_row <= row_1based <= merged.max_row
            and merged.min_col <= col_1based <= merged.max_col
        ):
            return str(merged)
    return ""


def merged_origin_0based(worksheet, row0: int, col0: int) -> Position:
    if worksheet is None:
        return (row0, col0)
    row1, col1 = row0 + 1, col0 + 1
    for merged in worksheet.merged_cells.ranges:
        if (
            merged.min_row <= row1 <= merged.max_row
            and merged.min_col <= col1 <= merged.max_col
        ):
            return (merged.min_row - 1, merged.min_col - 1)
    return (row0, col0)


def pair_image_positions(
    objects_a: Dict[Position, object],
    objects_b: Dict[Position, object],
    worksheet_a=None,
    worksheet_b=None,
    row_slop: int = 1,
    col_slop: int = 1,
    cancel_cb=None,
) -> List[Tuple[Optional[Position], Optional[Position]]]:
    """두 Excel의 이미지 앵커를 구조적으로 짝짓는다.

    픽셀이나 PASS/FAIL을 비교하지 않는다. 동일 셀, 같은 병합영역, 인접 셀
    순서로 위치만 맞춰 Double/Triple 화면의 한 행을 구성한다.
    """
    unused_a = set(objects_a)
    unused_b = set(objects_b)
    pairs: List[Tuple[Optional[Position], Optional[Position]]] = []

    def raise_if_cancelled() -> None:
        if cancel_cb is not None and cancel_cb():
            raise InterruptedError("사용자 취소")

    for position in sorted(unused_a & unused_b):
        raise_if_cancelled()
        pairs.append((position, position))
        unused_a.remove(position)
        unused_b.remove(position)

    origins_b: Dict[Position, List[Position]] = {}
    for position in unused_b:
        raise_if_cancelled()
        origins_b.setdefault(
            merged_origin_0based(worksheet_b, *position), []
        ).append(position)
    for positions in origins_b.values():
        positions.sort()

    for position_a in sorted(list(unused_a)):
        raise_if_cancelled()
        origin = merged_origin_0based(worksheet_a, *position_a)
        candidates = [p for p in origins_b.get(origin, []) if p in unused_b]
        if not candidates:
            continue
        position_b = min(
            candidates,
            key=lambda p: (
                abs(p[0] - position_a[0]) + abs(p[1] - position_a[1]),
                p,
            ),
        )
        pairs.append((position_a, position_b))
        unused_a.remove(position_a)
        unused_b.remove(position_b)

    row_slop = max(0, int(row_slop))
    col_slop = max(0, int(col_slop))
    candidates = []
    for position_a in unused_a:
        raise_if_cancelled()
        for position_b in unused_b:
            raise_if_cancelled()
            row_delta = position_b[0] - position_a[0]
            col_delta = position_b[1] - position_a[1]
            if abs(row_delta) <= row_slop and abs(col_delta) <= col_slop:
                priority = (abs(col_delta), abs(row_delta), row_delta, col_delta)
                candidates.append((priority, position_a, position_b))
    candidates.sort(key=lambda item: (item[0], item[1], item[2]))
    for _priority, position_a, position_b in candidates:
        raise_if_cancelled()
        if position_a not in unused_a or position_b not in unused_b:
            continue
        pairs.append((position_a, position_b))
        unused_a.remove(position_a)
        unused_b.remove(position_b)

    pairs.extend((position, None) for position in sorted(unused_a))
    pairs.extend((None, position) for position in sorted(unused_b))
    pairs.sort(
        key=lambda pair: pair[0]
        if pair[0] is not None
        else pair[1]
        if pair[1] is not None
        else (10**9, 10**9)
    )
    return pairs


def extract_image_objects_by_position(worksheet) -> Dict[Position, object]:
    objects: Dict[Position, object] = {}
    if worksheet is None:
        return objects
    for image_obj in getattr(worksheet, "_images", []):
        anchor = getattr(image_obj, "anchor", None)
        if not hasattr(anchor, "_from"):
            continue
        position = (anchor._from.row, anchor._from.col)
        if position in objects:
            cell = f"{get_column_letter(position[1] + 1)}{position[0] + 1}"
            raise ValueError(
                f"'{worksheet.title}' 시트의 {cell} 위치에 이미지가 2개 이상 있습니다. "
                "각 이미지를 서로 다른 셀 위치에 배치해 주세요."
            )
        objects[position] = image_obj
    return objects


def extract_image_meta(worksheet, sheet_index: int, image_obj) -> ExtractedImage:
    anchor = image_obj.anchor._from
    row1, col1 = anchor.row + 1, anchor.col + 1
    return ExtractedImage(
        sheet_index=sheet_index,
        sheet_name=worksheet.title,
        cell_address=f"{get_column_letter(col1)}{row1}",
        merged_range=find_merged_range(worksheet, row1, col1),
        anchor_row=anchor.row,
        anchor_col=anchor.col,
    )


def make_null_image(
    sheet_index: int,
    sheet_name: str,
    row: int,
    col: int,
) -> ExtractedImage:
    return ExtractedImage(
        sheet_index=sheet_index,
        sheet_name=sheet_name,
        cell_address=f"{get_column_letter(col + 1)}{row + 1}",
        merged_range="",
        anchor_row=row,
        anchor_col=col,
        is_null=True,
    )


def guess_image_extension(raw: bytes) -> str:
    if raw.startswith(b"\x89PNG\r\n\x1a\n"):
        return ".png"
    if raw[:3] == b"\xff\xd8\xff":
        return ".jpg"
    if raw[:6] in (b"GIF87a", b"GIF89a"):
        return ".gif"
    if raw[:2] == b"BM":
        return ".bmp"
    if len(raw) >= 12 and raw[:4] == b"RIFF" and raw[8:12] == b"WEBP":
        return ".webp"
    return ".bin"


def persist_source_from_obj(
    extracted: ExtractedImage,
    image_obj,
    preview_dir: str,
    side: str,
) -> bytes:
    raw = image_obj._data()
    extension = guess_image_extension(raw)
    filename = (
        f"s{extracted.sheet_index}_r{extracted.anchor_row}_c{extracted.anchor_col}"
        f"_{side}{extension}"
    )
    path = os.path.join(preview_dir, filename)
    with open(path, "wb") as file_handle:
        file_handle.write(raw)
    extracted.image = None
    extracted.source_path = path
    extracted.preview_path = path
    return raw


def load_preview_pil(path: Optional[str]) -> Optional[Image.Image]:
    if not path or not os.path.isfile(path):
        return None
    try:
        with Image.open(path) as image:
            return image.convert("RGB").copy()
    except (OSError, ValueError):
        return None


def load_extracted_pil(extracted: Optional[ExtractedImage]) -> Optional[Image.Image]:
    if extracted is None or extracted.is_null:
        return None
    for path in (extracted.source_path, extracted.preview_path):
        image = load_preview_pil(path)
        if image is not None:
            return image
    if extracted.image is not None:
        return extracted.image.convert("RGB")
    return None
