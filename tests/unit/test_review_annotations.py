"""불량/메모 자동저장과 Excel 추출 회귀 테스트."""

from pathlib import Path

from openpyxl import load_workbook

from models import ExtractedImage
from review_annotations import ReviewAnnotationStore, export_annotations_xlsx


def _image() -> ExtractedImage:
    return ExtractedImage(
        sheet_index=2,
        sheet_name="Power off_LCx",
        cell_address="D693",
        merged_range="D693:J712",
        anchor_row=692,
        anchor_col=3,
        source_image_id="sheet2-image7",
        anchor_occurrence=2,
        anchor_count=3,
    )


def test_memo_and_defect_are_independent_and_restore(tmp_path: Path):
    source = tmp_path / "reference.xlsx"
    source.write_bytes(b"source")
    storage = tmp_path / "annotations.json"
    store = ReviewAnnotationStore(
        {"ref": str(source)}, "double", storage_path=storage
    )
    image = _image()

    store.update("ref", image, note="불량은 아니지만 재확인 필요")
    annotation = store.annotation_for("ref", image)
    assert annotation is not None
    assert annotation.is_defect is False
    assert annotation.note == "불량은 아니지만 재확인 필요"
    assert store.memo_only_count == 1

    store.update("ref", image, is_defect=True)
    store.update("ref", image, is_defect=False)
    assert store.annotation_for("ref", image).note == "불량은 아니지만 재확인 필요"
    assert store.memo_only_count == 1

    restored = ReviewAnnotationStore(
        {"ref": str(source)}, "double", storage_path=storage
    )
    assert restored.annotation_for("ref", image).note == "불량은 아니지만 재확인 필요"

    restored.update("ref", image, note="")
    assert restored.records() == []


def test_export_contains_only_defect_or_memo_records(tmp_path: Path):
    source = tmp_path / "reference.xlsx"
    source.write_bytes(b"source")
    store = ReviewAnnotationStore(
        {"ref": str(source)},
        "double",
        storage_path=tmp_path / "annotations.json",
    )
    image = _image()
    store.update("ref", image, note="=1+1")

    output = tmp_path / "review.xlsx"
    assert export_annotations_xlsx(output, store.records()) == 1

    workbook = load_workbook(output, data_only=False)
    sheet = workbook["검토 기록"]
    assert [sheet.cell(1, column).value for column in range(1, 13)] == [
        "번호",
        "기록 유형",
        "불량 여부",
        "대상",
        "파일명",
        "파일 경로",
        "시트",
        "셀 위치",
        "병합 영역",
        "중복 순번",
        "메모",
        "기록 시간",
    ]
    assert sheet["B2"].value == "메모"
    assert sheet["C2"].value == "아니오"
    assert sheet["D2"].value == "Excel Ref"
    assert sheet["H2"].value == "D693"
    assert sheet["J2"].value == "2/3"
    assert sheet["K2"].value == "=1+1"
    assert sheet["K2"].data_type == "s"
