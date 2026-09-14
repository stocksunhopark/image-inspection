"""Double/Triple Excel loading for the side-by-side visual inspector."""

from __future__ import annotations

import shutil
import tempfile
from pathlib import Path
from typing import Dict, Tuple

import pytest
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from PIL import Image

from excel_manager import load_extracted_pil
import inspection_service as inspection_service_module
from inspection_service import load_inspection_images
from models import MISSING_IMAGE, MISSING_SHEET, ExtractedImage, InspectionItem
from sheet_mapping import SheetMappingGroup, build_sheet_mapping_from_paths
from ui.sheet_mapping_dialog import SheetMappingDialog


Color = Tuple[int, int, int]
WorkbookSpec = Dict[str, Dict[str, object]]


REF_COLORS: Dict[str, Color] = {
    "A2": (220, 20, 20),
    "D10": (20, 220, 20),
    "F20": (20, 20, 220),
    "J30": (180, 20, 180),
    "P60": (20, 180, 180),
    "SECOND!C3": (120, 80, 40),
}
A_COLORS: Dict[str, Color] = {
    "A2": (210, 40, 30),
    "D11": (30, 210, 40),
    "H22": (30, 40, 210),
    "L40": (210, 160, 30),
    "P60": (40, 170, 170),
    "SECOND!C3": (100, 90, 50),
}
C_COLORS: Dict[str, Color] = {
    "A2": (190, 60, 50),
    "D9": (50, 190, 60),
    "H21": (50, 60, 190),
    "J30": (160, 50, 160),
    "R70": (90, 90, 20),
    "SECOND!C3": (70, 110, 70),
}
B_COLORS: Dict[str, Color] = {
    "A2": (200, 50, 40),
    "D9": (40, 200, 50),
    "G21": (40, 50, 200),
    "J30": (170, 40, 170),
    "L41": (200, 150, 40),
    "N50": (80, 80, 80),
    "SECOND!C3": (80, 100, 60),
}


def _write_workbook(path: Path, spec: WorkbookSpec) -> Path:
    asset_dir = path.parent / f"{path.stem}-images"
    asset_dir.mkdir()
    workbook = Workbook()

    for sheet_number, (sheet_name, sheet_spec) in enumerate(spec.items()):
        worksheet = workbook.active if sheet_number == 0 else workbook.create_sheet()
        worksheet.title = sheet_name
        for merged_range in sheet_spec.get("merges", ()):  # type: ignore[union-attr]
            worksheet.merge_cells(str(merged_range))
        images = sheet_spec.get("images", {})  # type: ignore[union-attr]
        for cell, color in images.items():  # type: ignore[union-attr]
            image_path = asset_dir / f"{sheet_name}-{cell}.png"
            Image.new("RGB", (18, 12), color).save(image_path, format="PNG")
            worksheet.add_image(XLImage(str(image_path)), str(cell))

    workbook.save(path)
    workbook.close()
    return path


@pytest.fixture
def inspection_books(tmp_path: Path) -> Dict[str, Path]:
    common_merge = ["F20:H22"]
    ref = _write_workbook(
        tmp_path / "ref.xlsx",
        {
            "MAIN": {
                "merges": common_merge,
                # Deliberately not in coordinate order.
                "images": {
                    "P60": REF_COLORS["P60"],
                    "J30": REF_COLORS["J30"],
                    "F20": REF_COLORS["F20"],
                    "D10": REF_COLORS["D10"],
                    "A2": REF_COLORS["A2"],
                },
            },
            "SECOND": {"images": {"C3": REF_COLORS["SECOND!C3"]}},
        },
    )
    compare_a = _write_workbook(
        tmp_path / "a.xlsx",
        {
            "MAIN": {
                "merges": common_merge,
                "images": {
                    "L40": A_COLORS["L40"],
                    "H22": A_COLORS["H22"],
                    "P60": A_COLORS["P60"],
                    "D11": A_COLORS["D11"],
                    "A2": A_COLORS["A2"],
                },
            },
            "SECOND": {"images": {"C3": A_COLORS["SECOND!C3"]}},
        },
    )
    compare_b = _write_workbook(
        tmp_path / "b.xlsx",
        {
            "MAIN": {
                "merges": common_merge,
                "images": {
                    "N50": B_COLORS["N50"],
                    "L41": B_COLORS["L41"],
                    "J30": B_COLORS["J30"],
                    "G21": B_COLORS["G21"],
                    "D9": B_COLORS["D9"],
                    "A2": B_COLORS["A2"],
                },
            },
            "SECOND": {"images": {"C3": B_COLORS["SECOND!C3"]}},
        },
    )
    compare_c = _write_workbook(
        tmp_path / "c.xlsx",
        {
            "MAIN": {
                "merges": common_merge,
                "images": {
                    "R70": C_COLORS["R70"],
                    "J30": C_COLORS["J30"],
                    "H21": C_COLORS["H21"],
                    "D9": C_COLORS["D9"],
                    "A2": C_COLORS["A2"],
                },
            },
            "SECOND": {"images": {"C3": C_COLORS["SECOND!C3"]}},
        },
    )
    return {"ref": ref, "a": compare_a, "b": compare_b, "c": compare_c}


def _all_items(items_by_sheet: Dict[int, list[InspectionItem]]) -> list[InspectionItem]:
    return [
        item
        for sheet_index in sorted(items_by_sheet)
        for item in items_by_sheet[sheet_index]
    ]


def _assert_real_lazy_image(
    extracted: ExtractedImage,
    preview_dir: str,
    expected_color: Color,
) -> None:
    assert extracted.is_null is False
    assert extracted.image is None
    assert extracted.source_path
    assert extracted.preview_path == extracted.source_path
    source_path = Path(extracted.source_path).resolve()
    assert source_path.is_file()
    assert Path(preview_dir).resolve() in source_path.parents
    loaded = load_extracted_pil(extracted)
    assert loaded is not None
    assert loaded.mode == "RGB"
    assert loaded.getpixel((0, 0)) == expected_color


def _assert_null_image(extracted: ExtractedImage, expected_cell: str) -> None:
    assert extracted.is_null is True
    assert extracted.cell_address == expected_cell
    assert extracted.image is None
    assert extracted.source_path is None
    assert extracted.preview_path is None
    assert load_extracted_pil(extracted) is None


def test_double_matches_exact_adjacent_and_merged_positions_and_keeps_missing_sides(
    inspection_books: Dict[str, Path],
):
    items_by_sheet, preview_dir = load_inspection_images(
        str(inspection_books["ref"]), str(inspection_books["a"])
    )
    try:
        assert list(items_by_sheet) == [1, 2]
        main = items_by_sheet[1]
        assert [item.index for item in main] == [1, 2, 3, 4, 5, 6]
        assert [item.cell_address for item in main] == [
            "A2",
            "D10",
            "F20",
            "J30",
            "L40",
            "P60",
        ]

        # Exact anchor, one-row tolerance, then same merged block even though
        # the anchors differ by two rows and two columns.
        assert (main[0].image_ref.cell_address, main[0].image_a.cell_address) == (
            "A2",
            "A2",
        )
        assert (main[1].image_ref.cell_address, main[1].image_a.cell_address) == (
            "D10",
            "D11",
        )
        assert (main[2].image_ref.cell_address, main[2].image_a.cell_address) == (
            "F20",
            "H22",
        )
        assert main[2].image_ref.merged_range == "F20:H22"
        assert main[2].image_a.merged_range == "F20:H22"

        _assert_null_image(main[3].image_a, "J30")
        _assert_null_image(main[4].image_ref, "L40")
        assert main[3].image_ref.cell_address == "J30"
        assert main[4].image_a.cell_address == "L40"
        assert all(item.image_b is None for item in _all_items(items_by_sheet))
        assert all(item.image_c is None for item in _all_items(items_by_sheet))

        assert [item.cell_address for item in items_by_sheet[2]] == ["C3"]
        assert items_by_sheet[2][0].index == 1
    finally:
        shutil.rmtree(preview_dir, ignore_errors=True)


def test_triple_aligns_all_sides_persists_lazy_sources_and_reports_progress(
    inspection_books: Dict[str, Path],
):
    progress = []
    statuses = []
    items_by_sheet, preview_dir = load_inspection_images(
        str(inspection_books["ref"]),
        str(inspection_books["a"]),
        str(inspection_books["b"]),
        progress_cb=lambda current, total, cell: progress.append(
            (current, total, cell)
        ),
        status_cb=statuses.append,
    )
    try:
        main = items_by_sheet[1]
        assert [item.cell_address for item in main] == [
            "A2",
            "D10",
            "F20",
            "J30",
            "L40",
            "N50",
            "P60",
        ]
        assert [
            (
                item.image_ref.cell_address,
                item.image_a.cell_address,
                item.image_b.cell_address if item.image_b is not None else None,
            )
            for item in main[:3]
        ] == [
            ("A2", "A2", "A2"),
            ("D10", "D11", "D9"),
            ("F20", "H22", "G21"),
        ]

        # A is absent, Ref is absent, then both Ref/A are absent.
        assert main[3].image_b is not None
        _assert_null_image(main[3].image_a, "J30")
        _assert_null_image(main[4].image_ref, "L40")
        assert main[4].image_b is not None
        assert main[4].image_b.cell_address == "L41"
        _assert_null_image(main[5].image_ref, "N50")
        _assert_null_image(main[5].image_a, "N50")
        assert main[5].image_b is not None
        assert main[5].image_b.cell_address == "N50"
        assert main[6].image_b is not None
        _assert_null_image(main[6].image_b, "P60")

        first = main[0]
        assert first.image_b is not None
        _assert_real_lazy_image(first.image_ref, preview_dir, REF_COLORS["A2"])
        _assert_real_lazy_image(first.image_a, preview_dir, A_COLORS["A2"])
        _assert_real_lazy_image(first.image_b, preview_dir, B_COLORS["A2"])
        _assert_real_lazy_image(main[4].image_a, preview_dir, A_COLORS["L40"])
        assert main[4].image_b is not None
        _assert_real_lazy_image(main[4].image_b, preview_dir, B_COLORS["L41"])

        all_items = _all_items(items_by_sheet)
        assert len(all_items) == 8
        assert [current for current, _total, _cell in progress] == list(range(1, 9))
        assert {total for _current, total, _cell in progress} == {8}
        assert all(str(cell).strip() for _current, _total, cell in progress)
        assert statuses and all(str(message).strip() for message in statuses)

        # Sources have been copied lazily, so all input workbooks can be moved
        # after the service returns (important for Windows file handles).
        for role, workbook_path in inspection_books.items():
            if role == "c":
                continue
            moved = workbook_path.with_name(f"{role}-moved.xlsx")
            workbook_path.rename(moved)
            assert moved.is_file()
        assert load_extracted_pil(first.image_ref) is not None
    finally:
        shutil.rmtree(preview_dir, ignore_errors=True)


def test_quadra_aligns_c_side_and_keeps_c_only_rows(
    inspection_books: Dict[str, Path],
):
    items_by_sheet, preview_dir = load_inspection_images(
        str(inspection_books["ref"]),
        str(inspection_books["a"]),
        str(inspection_books["b"]),
        str(inspection_books["c"]),
    )
    try:
        main = items_by_sheet[1]
        assert [item.cell_address for item in main] == [
            "A2",
            "D10",
            "F20",
            "J30",
            "L40",
            "N50",
            "P60",
            "R70",
        ]
        first = main[0]
        assert first.image_c is not None
        _assert_real_lazy_image(first.image_ref, preview_dir, REF_COLORS["A2"])
        _assert_real_lazy_image(first.image_a, preview_dir, A_COLORS["A2"])
        _assert_real_lazy_image(first.image_b, preview_dir, B_COLORS["A2"])
        _assert_real_lazy_image(first.image_c, preview_dir, C_COLORS["A2"])
        assert first.image_c.cell_address == "A2"
        assert main[1].image_c is not None
        assert main[1].image_c.cell_address == "D9"
        assert main[2].image_c is not None
        assert main[2].image_c.cell_address == "H21"
        assert main[3].image_c is not None
        assert main[3].image_c.cell_address == "J30"
        last = main[7]
        assert last.image_c is not None
        assert last.image_c.cell_address == "R70"
        _assert_null_image(last.image_ref, "R70")
        _assert_null_image(last.image_a, "R70")
        _assert_null_image(last.image_b, "R70")
        second = items_by_sheet[2][0]
        assert second.image_c is not None
        _assert_real_lazy_image(second.image_c, preview_dir, C_COLORS["SECOND!C3"])
    finally:
        shutil.rmtree(preview_dir, ignore_errors=True)


def test_quadra_requires_compare_b(tmp_path: Path):
    ref = _write_workbook(
        tmp_path / "ref-c-only.xlsx",
        {"MAIN": {"images": {"A1": (10, 20, 30)}}},
    )
    compare_a = _write_workbook(
        tmp_path / "a-c-only.xlsx",
        {"MAIN": {"images": {"A1": (10, 20, 30)}}},
    )
    compare_c = _write_workbook(
        tmp_path / "c-only.xlsx",
        {"MAIN": {"images": {"A1": (10, 20, 30)}}},
    )
    with pytest.raises(ValueError, match="비교B"):
        load_inspection_images(str(ref), str(compare_a), path_c=str(compare_c))


def test_cancellation_raises_and_removes_partial_preview_directory(
    inspection_books: Dict[str, Path],
    tmp_path: Path,
    monkeypatch,
):
    service_temp_root = tmp_path / "service-temp"
    service_temp_root.mkdir()
    monkeypatch.setattr(tempfile, "tempdir", str(service_temp_root))
    progress = []

    with pytest.raises(InterruptedError, match="취소"):
        load_inspection_images(
            str(inspection_books["ref"]),
            str(inspection_books["a"]),
            str(inspection_books["b"]),
            progress_cb=lambda current, total, cell: progress.append(
                (current, total, cell)
            ),
            cancel_cb=lambda: bool(progress),
        )

    assert len(progress) == 1
    assert progress[0][:2] == (1, 8)
    assert list(service_temp_root.iterdir()) == []


def test_mismatched_sheet_tabs_form_union_and_keep_empty_sheets(tmp_path: Path):
    ref = _write_workbook(
        tmp_path / "ref-mismatch.xlsx",
        {"MAIN": {"images": {"A1": (10, 20, 30)}}, "SECOND": {"images": {}}},
    )
    compare = _write_workbook(
        tmp_path / "a-mismatch.xlsx",
        {"MAIN": {"images": {"A1": (10, 20, 30)}}, "OTHER": {"images": {}}},
    )

    result = load_inspection_images(str(ref), str(compare))
    try:
        assert list(result.items_by_sheet) == [1, 2, 3]
        assert [
            result.sheet_infos[index].display_name for index in result.sheet_infos
        ] == ["MAIN", "SECOND", "OTHER"]
        assert [
            result.sheet_infos[index].status_text for index in result.sheet_infos
        ] == ["REF + A", "REF ONLY", "A ONLY"]

        second = result.items_by_sheet[2][0]
        assert second.is_empty_sheet is True
        assert second.image_ref.missing_reason == MISSING_IMAGE
        assert second.image_a.missing_reason == MISSING_SHEET
        assert second.image_ref.cell_address == "-"
        assert second.image_a.cell_address == "-"

        other = result.items_by_sheet[3][0]
        assert other.is_empty_sheet is True
        assert other.image_ref.missing_reason == MISSING_SHEET
        assert other.image_a.missing_reason == MISSING_IMAGE

        assert result.integrity_report.review_sheet_count == 3
        assert result.integrity_report.source_image_count == 2
        assert result.integrity_report.queue_image_count == 2
    finally:
        shutil.rmtree(result.preview_dir, ignore_errors=True)


def test_custom_sheet_mapping_groups_different_names_and_controls_order(
    tmp_path: Path,
):
    ref = _write_workbook(
        tmp_path / "ref-custom.xlsx",
        {
            "MAIN": {"images": {"A1": (10, 20, 30)}},
            "CLK": {"images": {"B2": (40, 50, 60)}},
        },
    )
    compare = _write_workbook(
        tmp_path / "a-custom.xlsx",
        {
            "MAIN": {"images": {"A1": (70, 80, 90)}},
            "Clock": {"images": {"B2": (100, 110, 120)}},
        },
    )
    paths = {"ref": str(ref), "a": str(compare)}
    plan = build_sheet_mapping_from_paths(paths, ("ref", "a"))
    source_ids = {
        (source.role, source.name): source.source_sheet_id
        for source in plan.sources
    }
    plan.groups = [
        SheetMappingGroup(
            display_name="CLK 사용자 비교",
            role_sheet_ids={
                "ref": source_ids[("ref", "CLK")],
                "a": source_ids[("a", "Clock")],
            },
        ),
        SheetMappingGroup(
            display_name="MAIN",
            role_sheet_ids={
                "ref": source_ids[("ref", "MAIN")],
                "a": source_ids[("a", "MAIN")],
            },
        ),
    ]

    result = load_inspection_images(
        str(ref), str(compare), sheet_mapping=plan
    )
    try:
        assert [info.display_name for info in result.sheet_infos.values()] == [
            "CLK 사용자 비교",
            "MAIN",
        ]
        assert result.sheet_infos[1].role_sheet_names == {
            "ref": "CLK",
            "a": "Clock",
        }
        assert result.sheet_infos[1].status_text == "REF + A"
        assert result.items_by_sheet[1][0].image_ref.sheet_name == "CLK"
        assert result.items_by_sheet[1][0].image_a.sheet_name == "Clock"
        assert result.integrity_report.source_image_count == 4
        assert result.integrity_report.queue_image_count == 4
    finally:
        shutil.rmtree(result.preview_dir, ignore_errors=True)


@pytest.mark.parametrize(
    "roles",
    [
        ("ref", "a"),
        ("ref", "a", "b"),
        ("ref", "a", "b", "c"),
    ],
    ids=["double", "triple", "quadra"],
)
def test_custom_mapping_loads_real_embedded_images_in_every_mode(
    tmp_path: Path, roles: Tuple[str, ...], qapp
):
    colors = {
        "ref": (210, 20, 20),
        "a": (20, 210, 20),
        "b": (20, 20, 210),
        "c": (210, 210, 20),
    }
    sheet_names = {
        "ref": "REF_DIFFERENT",
        "a": "A_DIFFERENT",
        "b": "B_DIFFERENT",
        "c": "C_DIFFERENT",
    }
    paths = {
        role: str(
            _write_workbook(
                tmp_path / f"{role}.xlsx",
                {
                    sheet_names[role]: {
                        "images": {"A1": colors[role]}
                    }
                },
            )
        )
        for role in roles
    }
    automatic = build_sheet_mapping_from_paths(paths, roles)
    source_ids = {
        source.role: source.source_sheet_id for source in automatic.sources
    }
    original_groups = [
        (group.display_name, dict(group.role_sheet_ids))
        for group in automatic.groups
    ]
    dialog = SheetMappingDialog(automatic, automatic)
    try:
        dialog._add_group()
        custom_group = dialog._plan.groups[-1]
        for role in roles:
            dialog._assign_source(custom_group, role, source_ids[role])

        assert custom_group.display_name == "새 비교 그룹 1"
        assert [
            (group.display_name, dict(group.role_sheet_ids))
            for group in dialog._plan.groups[:-1]
        ] == original_groups

        dialog._accept_mapping()
        assert dialog.result_plan is not None
        plan = dialog.result_plan
    finally:
        dialog.close()
        dialog.deleteLater()

    positional_paths = [paths[role] for role in roles]

    result = load_inspection_images(
        *positional_paths, sheet_mapping=plan
    )
    try:
        assert list(result.items_by_sheet) == [1]
        assert result.sheet_infos[1].display_name == "새 비교 그룹 1"
        assert result.sheet_infos[1].role_sheet_names == {
            role: sheet_names[role] for role in roles
        }
        item = result.items_by_sheet[1][0]
        assert tuple(role for role, _image in item.side_images()) == roles
        for role, extracted in item.side_images():
            assert extracted.is_null is False
            assert extracted.sheet_name == sheet_names[role]
            image = load_extracted_pil(extracted)
            assert image is not None
            assert image.getpixel((0, 0)) == colors[role]
        assert result.integrity_report.source_image_count == len(roles)
        assert result.integrity_report.queue_image_count == len(roles)
    finally:
        shutil.rmtree(result.preview_dir, ignore_errors=True)


@pytest.mark.parametrize(
    "roles",
    [
        ("ref", "a"),
        ("ref", "a", "b"),
        ("ref", "a", "b", "c"),
    ],
    ids=["double", "triple", "quadra"],
)
def test_inactive_mapping_group_skips_real_images_in_every_mode(
    tmp_path: Path, roles: Tuple[str, ...]
):
    active_colors = {
        "ref": (201, 11, 11),
        "a": (11, 201, 11),
        "b": (11, 11, 201),
        "c": (201, 201, 11),
    }
    paths = {
        role: str(
            _write_workbook(
                tmp_path / f"{role}-activation.xlsx",
                {
                    "ACTIVE": {"images": {"A1": active_colors[role]}},
                    "SKIPPED": {"images": {"B2": (90, 90, 90)}},
                },
            )
        )
        for role in roles
    }
    plan = build_sheet_mapping_from_paths(paths, roles)
    assert [group.display_name for group in plan.groups] == [
        "ACTIVE",
        "SKIPPED",
    ]
    plan.groups[1].enabled = False

    result = load_inspection_images(
        *(paths[role] for role in roles),
        sheet_mapping=plan,
    )
    try:
        assert list(result.items_by_sheet) == [1]
        assert list(result.sheet_infos) == [1]
        assert result.sheet_infos[1].display_name == "ACTIVE"
        item = result.items_by_sheet[1][0]
        for role, extracted in item.side_images():
            image = load_extracted_pil(extracted)
            assert image is not None
            assert extracted.sheet_name == "ACTIVE"
            assert image.getpixel((0, 0)) == active_colors[role]
        assert result.integrity_report.role_sheet_counts == tuple(
            (role, 1) for role in roles
        )
        assert result.integrity_report.review_sheet_count == 1
        assert result.integrity_report.source_image_count == len(roles)
        assert result.integrity_report.queue_image_count == len(roles)
    finally:
        shutil.rmtree(result.preview_dir, ignore_errors=True)


def test_triple_sheet_union_uses_ref_order_name_matching_and_role_statuses(
    tmp_path: Path,
):
    ref = _write_workbook(
        tmp_path / "ref-union.xlsx",
        {
            "ROOM": {"images": {"A1": (220, 20, 20)}},
            "REF_ONLY": {"images": {"B2": (20, 220, 20)}},
            "EMPTY": {"images": {}},
        },
    )
    compare_a = _write_workbook(
        tmp_path / "a-union.xlsx",
        {
            "A_ONLY": {"images": {"A3": (20, 20, 220)}},
            "room": {"images": {"A1": (200, 40, 40)}},
            "EMPTY": {"images": {}},
        },
    )
    compare_b = _write_workbook(
        tmp_path / "b-union.xlsx",
        {
            "B_ONLY": {"images": {"C4": (180, 20, 180)}},
            "ROOM": {"images": {"A1": (180, 50, 50)}},
            "REF_ONLY": {"images": {"B2": (50, 180, 50)}},
        },
    )

    result = load_inspection_images(str(ref), str(compare_a), str(compare_b))
    try:
        assert [info.display_name for info in result.sheet_infos.values()] == [
            "ROOM",
            "REF_ONLY",
            "EMPTY",
            "A_ONLY",
            "B_ONLY",
        ]
        assert [info.status_text for info in result.sheet_infos.values()] == [
            "REF + A + B",
            "REF + B",
            "REF + A",
            "A ONLY",
            "B ONLY",
        ]

        room = result.items_by_sheet[1][0]
        assert room.image_ref.sheet_name == "ROOM"
        assert room.image_a.sheet_name == "room"
        assert room.image_b is not None
        assert room.image_b.sheet_name == "ROOM"

        ref_only = result.items_by_sheet[2][0]
        assert ref_only.image_a.missing_reason == MISSING_SHEET
        assert ref_only.image_b is not None
        assert ref_only.image_b.is_null is False

        empty = result.items_by_sheet[3][0]
        assert empty.is_empty_sheet is True
        assert empty.image_ref.missing_reason == MISSING_IMAGE
        assert empty.image_a.missing_reason == MISSING_IMAGE
        assert empty.image_b is not None
        assert empty.image_b.missing_reason == MISSING_SHEET

        assert result.integrity_report.role_sheet_counts == (
            ("ref", 3),
            ("a", 3),
            ("b", 3),
        )
        assert result.integrity_report.review_sheet_count == 5
        assert result.integrity_report.source_image_count == 7
        assert result.integrity_report.queue_image_count == 7
        source_ids = [
            image.source_image_id
            for item in _all_items(result.items_by_sheet)
            for _role, image in item.source_images()
        ]
        assert len(source_ids) == len(set(source_ids)) == 7
    finally:
        shutil.rmtree(result.preview_dir, ignore_errors=True)


def test_quadra_sheet_union_keeps_c_only_sheet_and_missing_roles(tmp_path: Path):
    ref = _write_workbook(
        tmp_path / "ref-quad-union.xlsx",
        {"ROOM": {"images": {"A1": (220, 20, 20)}}},
    )
    compare_a = _write_workbook(
        tmp_path / "a-quad-union.xlsx",
        {"A_ONLY": {"images": {"B2": (20, 220, 20)}}},
    )
    compare_b = _write_workbook(
        tmp_path / "b-quad-union.xlsx",
        {"ROOM": {"images": {"A1": (20, 20, 220)}}},
    )
    compare_c = _write_workbook(
        tmp_path / "c-quad-union.xlsx",
        {"C_ONLY": {"images": {"C3": (180, 20, 180)}}},
    )

    result = load_inspection_images(
        str(ref), str(compare_a), str(compare_b), str(compare_c)
    )
    try:
        assert [info.display_name for info in result.sheet_infos.values()] == [
            "ROOM",
            "A_ONLY",
            "C_ONLY",
        ]
        assert [info.status_text for info in result.sheet_infos.values()] == [
            "REF + B",
            "A ONLY",
            "C ONLY",
        ]
        c_only = result.items_by_sheet[3][0]
        assert c_only.image_ref.missing_reason == MISSING_SHEET
        assert c_only.image_a.missing_reason == MISSING_SHEET
        assert c_only.image_b is not None
        assert c_only.image_b.missing_reason == MISSING_SHEET
        assert c_only.image_c is not None
        assert c_only.image_c.is_null is False
        assert result.integrity_report.source_image_count == 4
    finally:
        shutil.rmtree(result.preview_dir, ignore_errors=True)


def test_ambiguous_normalized_sheet_names_stay_separate_with_warning(
    tmp_path: Path,
):
    ref = _write_workbook(
        tmp_path / "ref-ambiguous.xlsx",
        {
            "ROOM": {"images": {"A1": (220, 20, 20)}},
            "ROOM ": {"images": {"B2": (20, 220, 20)}},
        },
    )
    compare = _write_workbook(
        tmp_path / "a-ambiguous.xlsx",
        {"room": {"images": {"C3": (20, 20, 220)}}},
    )

    result = load_inspection_images(str(ref), str(compare))
    try:
        assert len(result.sheet_infos) == 3
        assert [info.status_text for info in result.sheet_infos.values()] == [
            "REF ONLY",
            "REF ONLY",
            "A ONLY",
        ]
        assert len(result.warnings) == 1
        assert "모호" in result.warnings[0]
        assert result.integrity_report.source_image_count == 3
    finally:
        shutil.rmtree(result.preview_dir, ignore_errors=True)


def test_integrity_check_rejects_duplicate_queue_source_image(
    tmp_path: Path,
    monkeypatch,
):
    ref = _write_workbook(
        tmp_path / "ref-integrity.xlsx",
        {"MAIN": {"images": {"A1": (220, 20, 20)}}},
    )
    compare = _write_workbook(
        tmp_path / "a-integrity.xlsx",
        {"MAIN": {"images": {"A1": (20, 220, 20)}}},
    )
    original_align = inspection_service_module._align_positions

    def duplicate_first_row(*args, **kwargs):
        rows = original_align(*args, **kwargs)
        return rows + [rows[0]]

    monkeypatch.setattr(
        inspection_service_module, "_align_positions", duplicate_first_row
    )
    with pytest.raises(ValueError, match="중복/추가 Source Image ID"):
        load_inspection_images(str(ref), str(compare))


def test_duplicate_image_anchor_is_rejected_instead_of_dropping_an_image(
    tmp_path: Path,
):
    source_a = tmp_path / "duplicate-a.png"
    source_b = tmp_path / "duplicate-b.png"
    Image.new("RGB", (12, 8), (255, 0, 0)).save(source_a)
    Image.new("RGB", (12, 8), (0, 255, 0)).save(source_b)
    ref_book = Workbook()
    ref_sheet = ref_book.active
    ref_sheet.title = "MAIN"
    ref_sheet.add_image(XLImage(str(source_a)), "A1")
    ref_sheet.add_image(XLImage(str(source_b)), "A1")
    ref = tmp_path / "ref-duplicate.xlsx"
    ref_book.save(ref)
    ref_book.close()

    compare = _write_workbook(
        tmp_path / "a-duplicate.xlsx",
        {"MAIN": {"images": {"A1": (10, 20, 30)}}},
    )

    with pytest.raises(ValueError, match="이미지가 2개 이상"):
        load_inspection_images(str(ref), str(compare))
