"""xlsx 시트 이름만 workbook.xml 에서 읽는 경로 테스트."""
from openpyxl import Workbook

from excel_manager import read_sheet_index_names, _sheet_names_from_workbook_xml


def test_read_sheet_index_names_matches_tab_order(tmp_path):
    path = tmp_path / "sheets.xlsx"
    wb = Workbook()
    wb.active.title = "CLK"
    wb.create_sheet("ROOM")
    wb.create_sheet("HOT")
    wb.save(path)

    assert read_sheet_index_names(str(path)) == [
        (1, "CLK"),
        (2, "ROOM"),
        (3, "HOT"),
    ]
    assert _sheet_names_from_workbook_xml(str(path)) == ["CLK", "ROOM", "HOT"]


def test_sheet_names_xml_returns_none_for_missing_file(tmp_path):
    assert _sheet_names_from_workbook_xml(str(tmp_path / "nope.xlsx")) is None
